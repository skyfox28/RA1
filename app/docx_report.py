import io
import os

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

RED = "C00000"
RED_BG = "FCE4E4"
HEADER_BG = "D9E2F3"

DRIFT_THRESHOLD_PCT = 10


def _shade_cell(cell, color_hex):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color_hex)
    tc_pr.append(shd)


def _set_cell_text(cell, text, bold=False, color=None, size=8):
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run("" if text is None else str(text))
    run.bold = bold
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def _fmt(v):
    if v is None:
        return "—"
    return v


def _fmt_date(v):
    return v.strftime("%d/%m/%Y") if v else "—"


def _make_landscape(document):
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width


def build_report_docx(report, upload_folder):
    doc = Document()
    _make_landscape(doc)

    title = doc.add_heading(f"Compte-rendu de contrôle journalier – Zone {report.zone}", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    info = doc.add_paragraph()
    info.add_run("Date du contrôle : ").bold = True
    info.add_run(f"{report.report_date.strftime('%d/%m/%Y')}\n")
    info.add_run("Responsable du contrôle : ").bold = True
    info.add_run(f"{report.author or '—'}\n")

    if report.last_sap_import:
        li = report.last_sap_import
        info.add_run("Extraction stock (LX02) : ").bold = True
        info.add_run(f"{li.original_filename} — {len(li.rows)} lignes, importée le "
                     f"{li.imported_at.strftime('%d/%m/%Y %H:%M')}\n")
    if report.last_movement_import:
        lm = report.last_movement_import
        info.add_run("Extraction mouvements (LT27) : ").bold = True
        info.add_run(f"{lm.original_filename} — {len(lm.rows)} lignes, importée le "
                     f"{lm.imported_at.strftime('%d/%m/%Y %H:%M')}\n")
    if report.notes:
        info.add_run("Notes générales : ").bold = True
        info.add_run(report.notes)

    # --- Synthèse ---------------------------------------------------------
    doc.add_heading("Synthèse", level=1)
    total = report.total_lines
    drifts = len(report.drift_lines)
    rate = report.drift_rate

    summary = doc.add_paragraph()
    summary.add_run("Nombre de contrôles réalisés : ").bold = True
    summary.add_run(f"{total}\n")
    summary.add_run("Nombre d'écarts / dérives constatés : ").bold = True
    run = summary.add_run(f"{drifts} ({rate}%)\n")
    if drifts:
        run.bold = True
        run.font.color.rgb = RGBColor.from_string(RED)

    if total and rate >= DRIFT_THRESHOLD_PCT:
        alert = doc.add_paragraph()
        alert_run = alert.add_run(
            "⚠ Taux de dérive élevé sur cette zone — une analyse des causes et un plan "
            "d'action sont recommandés."
        )
        alert_run.bold = True
        alert_run.font.color.rgb = RGBColor.from_string(RED)
    elif total:
        doc.add_paragraph("Aucune dérive significative détectée sur cette zone ce jour.")

    # --- Détail des contrôles ---------------------------------------------
    doc.add_heading("Détail des contrôles physiques et informatiques", level=1)
    if report.lines:
        headers = [
            "Article", "Désignation", "Lot", "Unité de stock", "UM",
            "Qté SAP", "Qté physique", "Écart", "DLC", "Type de contrôle",
            "Personne", "Action réalisée", "Statut",
        ]
        table = doc.add_table(rows=1, cols=len(headers))
        table.style = "Light Grid Accent 1"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, h in enumerate(headers):
            _set_cell_text(table.rows[0].cells[i], h, bold=True)
            _shade_cell(table.rows[0].cells[i], HEADER_BG)

        for line in report.lines:
            row = table.add_row()
            values = [
                line.article, line.designation, line.lot, line.unite_stock, line.um,
                line.qte_sap, line.qte_physique, line.ecart, _fmt_date(line.dlc),
                line.type_controle, line.personne, line.action, line.statut,
            ]
            drift = line.is_drift
            for i, val in enumerate(values):
                _set_cell_text(row.cells[i], _fmt(val), color=RED if drift else None)
            if drift:
                for cell in row.cells:
                    _shade_cell(cell, RED_BG)
    else:
        doc.add_paragraph("Aucune ligne de contrôle enregistrée.")

    # --- Photos générales ---------------------------------------------------
    if report.photos:
        doc.add_heading("Photos du contrôle physique de la zone", level=1)
        for photo in report.photos:
            _add_photo(doc, upload_folder, photo)

    lines_with_photos = [l for l in report.lines if l.photos]
    if lines_with_photos:
        doc.add_heading("Photos associées aux contrôles", level=1)
        for line in lines_with_photos:
            caption = " — ".join(p for p in [line.article, line.designation] if p) or "Contrôle"
            doc.add_paragraph(caption, style="Intense Quote")
            for photo in line.photos:
                _add_photo(doc, upload_folder, photo)

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer


def _add_photo(doc, upload_folder, photo):
    abs_path = os.path.join(upload_folder, photo.filename)
    if os.path.exists(abs_path):
        try:
            doc.add_picture(abs_path, width=Inches(4))
        except Exception:
            doc.add_paragraph(f"[Image illisible : {photo.caption}]")
    if photo.caption:
        cap = doc.add_paragraph(photo.caption)
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in cap.runs:
            run.italic = True
            run.font.size = Pt(8)
