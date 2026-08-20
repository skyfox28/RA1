"""Lecture flexible des extractions Excel SAP (LX02 - stock, LT27 - ordres de
transfert / mouvements d'unités de manutention).

Les extractions SAP n'ont pas toujours exactement les mêmes intitulés de
colonnes selon la configuration du poste. On détecte donc les colonnes par
mots-clés (accents ignorés), et l'utilisateur peut corriger le mapping à
l'écran avant de valider l'import.
"""
import unicodedata
from datetime import date as date_cls
from datetime import datetime

import openpyxl
from dateutil import parser as dateparser

# Rôles reconnus pour une extraction de stock (transaction LX02)
ROLE_KEYWORDS_LX02 = {
    "article": ["article"],
    "designation": ["designation"],
    "unite_stock": ["unite de stock", "unite stock"],
    "type_stock": ["type de stock", "type stock"],
    "emplacement": ["emplacement", "magasin"],
    "quantite": ["stock disponible", "quantite", "qte disponible", "quantity"],
    "um": ["unite de qte", "unite de mesure", "uom"],
    "lot": ["lot", "batch"],
    "dlc": ["peremption", "dlc", "expiration"],
    "date_dernier_mouvement": ["dernier mouvement", "date mouvement"],
}

# Rôles reconnus pour une extraction d'ordres de transfert (transaction LT27)
# suivi des unités de manutention / unités de stock (UM) et des mouvements
# informatiques (qui a réalisé le mouvement, d'où vers où, quand).
ROLE_KEYWORDS_LT27 = {
    "numero_ot": ["ordre de transfert"],
    "code_mouvement": ["code mouvement"],
    "magasin": ["magasin"],
    "article": ["article"],
    "lot": ["lot"],
    "utilisateur": ["utilisateur"],
    "quantite": ["qte theor", "quantite theor", "quantite"],
    "um": ["unite de quantite alternative", "unite de mesure", "uom"],
    "statut_ot": ["statut confirmation", "statut"],
    "type_stock": ["type de stock"],
    "type_magasin_origine": ["type magasin cedant"],
    "emplacement_origine": ["emplacement cedant"],
    "unite_stock_origine": ["unite stock cedant", "unite de stock cedant"],
    "type_magasin_destination": ["type magasin prenant"],
    "emplacement_destination": ["emplacement prenant"],
    "unite_stock_destination": ["unite de stock pren", "unite stock pren"],
    "date_mouvement": ["date confirmation"],
    "heure_mouvement": ["heure confirmation"],
}

ROLE_LABELS = {
    "article": "Article",
    "designation": "Désignation",
    "unite_stock": "Unité de stock",
    "type_stock": "Type de stock",
    "emplacement": "Emplacement / Zone",
    "quantite": "Quantité",
    "um": "Unité de mesure",
    "lot": "Lot",
    "dlc": "Date péremption / DLC",
    "date_dernier_mouvement": "Date dernier mouvement",
    "numero_ot": "N° Ordre de transfert",
    "code_mouvement": "Code mouvement (WM)",
    "magasin": "Magasin",
    "utilisateur": "Utilisateur SAP (a réalisé le mouvement)",
    "statut_ot": "Statut confirmation",
    "type_magasin_origine": "Type magasin cédant",
    "emplacement_origine": "Emplacement cédant (origine)",
    "unite_stock_origine": "Unité de stock cédant",
    "type_magasin_destination": "Type magasin prenant",
    "emplacement_destination": "Emplacement prenant (destination)",
    "unite_stock_destination": "Unité de stock prenant",
    "date_mouvement": "Date du mouvement",
    "heure_mouvement": "Heure du mouvement",
}

ROLE_KEYWORDS_BY_KIND = {"lx02": ROLE_KEYWORDS_LX02, "lt27": ROLE_KEYWORDS_LT27}


def _norm(s):
    if s is None:
        return ""
    s = str(s).strip().lower()
    s = "".join(c for c in unicodedata.normalize("NFD", s) if unicodedata.category(c) != "Mn")
    return s


def to_date(value):
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date_cls):
        return value
    try:
        return dateparser.parse(str(value), dayfirst=True).date()
    except (ValueError, TypeError, OverflowError):
        return None


def sniff_headers(filepath):
    wb = openpyxl.load_workbook(filepath, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    headers = []
    for cell in next(ws.iter_rows(min_row=1, max_row=1)):
        headers.append(str(cell.value).strip() if cell.value is not None else "")
    wb.close()
    return [h for h in headers if h]


def guess_mapping(headers, kind):
    role_keywords = ROLE_KEYWORDS_BY_KIND[kind]
    normed = [(h, _norm(h)) for h in headers]
    mapping = {}
    for role, keywords in role_keywords.items():
        best = None
        for h, nh in normed:
            if any(kw in nh for kw in keywords):
                best = h
                break
        mapping[role] = best
    return mapping


def preview_rows(filepath, limit=5):
    headers = sniff_headers(filepath)
    wb = openpyxl.load_workbook(filepath, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    rows = []
    for row in ws.iter_rows(min_row=2, max_row=1 + limit, values_only=True):
        rows.append(dict(zip(headers, row)))
    wb.close()
    return rows


def parse_rows(filepath, mapping):
    """mapping: role -> nom de colonne (ou None si non renseigné)."""
    headers = sniff_headers(filepath)
    col_index = {h: idx for idx, h in enumerate(headers)}

    role_index = {role: col_index[header] for role, header in mapping.items()
                  if header and header in col_index}

    wb = openpyxl.load_workbook(filepath, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    results = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row is None or all(v is None for v in row):
            continue
        item = {role: (row[idx] if idx < len(row) else None) for role, idx in role_index.items()}
        if not any(v not in (None, "") for v in item.values()):
            continue
        results.append(item)
    wb.close()
    return results
